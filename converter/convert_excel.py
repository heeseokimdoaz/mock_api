"""엑셀 파일을 TapaCross API 명세 형식의 JSON으로 변환"""

import json
import os
from collections import defaultdict
from pathlib import Path

import openpyxl

BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = BASE_DIR / "data" / "converted"

POLARITY_MAP = {"긍정": "1", "중립": "0", "부정": "2"}

SITE_TYPE_MAP = {
    "트위터": "twitter",
    "매스미디어": "media",
    "비디오": "youtube",
    "커뮤니티": "comm",
}

NEWSPAPER_REGION_MAP = {
    "조선일보": "national",
    "중앙일보": "national",
    "서울신문": "seoul",
    "국민일보": "seoul",
    "전북도민일보": "jeolla",
    "무등일보": "jeolla",
    "부산일보": "gyeongsang",
    "매일신문": "gyeongsang",
}

# 조선일보만 9컬럼(작성자 포함), 나머지는 8컬럼
NEWSPAPER_HAS_AUTHOR = {"조선일보"}


def convert_date(date_val) -> str:
    """'2026-02-12 23:20:29' -> '20260212232029'"""
    s = str(date_val).strip()
    return s.replace("-", "").replace(" ", "").replace(":", "")


def safe_str(val) -> str:
    if val is None:
        return ""
    return str(val).strip()


def load_twitter(filepath: str) -> list[dict]:
    """X(트위터).xlsx: header at row 6, data from row 7"""
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active
    docs = []
    for row in ws.iter_rows(min_row=7, max_row=ws.max_row, values_only=True):
        if not row[4]:
            continue
        docs.append({
            "create_date": convert_date(row[4]),
            "site_type": "twitter",
            "site_name": safe_str(row[1]),
            "title": None,
            "content": safe_str(row[3]),
            "url": safe_str(row[5]),
            "polarity": POLARITY_MAP.get(safe_str(row[6]), "0"),
        })
    return docs


def load_media(filepath: str) -> list[dict]:
    """매스미디어.xlsx: header at row 9, data from row 10
    Cols: 매체, 제목, 본문, 작성일, URL, 사이트 구분, 사이트명, 긍부정
    """
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active
    docs = []
    for row in ws.iter_rows(min_row=10, max_row=ws.max_row, values_only=True):
        if not row[3]:
            continue
        docs.append({
            "create_date": convert_date(row[3]),
            "site_type": "media",
            "site_name": safe_str(row[6]),
            "title": safe_str(row[1]) or None,
            "content": safe_str(row[2]),
            "url": safe_str(row[4]),
            "polarity": POLARITY_MAP.get(safe_str(row[7]), "0"),
        })
    return docs


def load_youtube(filepath: str) -> list[dict]:
    """유튜브 스크립트.xlsx: header at row 3, data from row 4
    Cols: 매체, 사이트명, 제목, 본문, 스크립트, 작성일, URL, 긍부정
    """
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active
    docs = []
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=True):
        if not row[5]:
            continue
        content = safe_str(row[3])
        script = safe_str(row[4])
        if script:
            content = content + "\n\n[스크립트]\n" + script
        docs.append({
            "create_date": convert_date(row[5]),
            "site_type": "youtube",
            "site_name": safe_str(row[1]),
            "title": safe_str(row[2]) or None,
            "content": content,
            "url": safe_str(row[6]),
            "polarity": POLARITY_MAP.get(safe_str(row[7]), "0"),
        })
    return docs


def load_community(filepath: str) -> list[dict]:
    """커뮤니티.xlsx: header at row 3, data from row 4
    Cols: 매체, 사이트명, 제목, 본문, 작성일, URL, 세부 구분, 긍부정
    """
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active
    docs = []
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=True):
        if not row[4]:
            continue
        docs.append({
            "create_date": convert_date(row[4]),
            "site_type": "comm",
            "site_name": safe_str(row[1]),
            "title": safe_str(row[2]) or None,
            "content": safe_str(row[3]),
            "url": safe_str(row[5]),
            "polarity": POLARITY_MAP.get(safe_str(row[7]), "0"),
        })
    return docs


def load_newspaper(filepath: str, paper_name: str) -> list[dict]:
    """신문 엑셀: header at row 2, data from row 3
    Schema B (8컬럼): 매체, 제목, 본문, 작성일, URL, 사이트 구분, 사이트명, 긍부정
    Schema A (9컬럼, 조선일보): 매체, 제목, 본문, 작성자, 작성일, URL, 사이트 구분, 사이트명, 긍부정
    """
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active
    has_author = paper_name in NEWSPAPER_HAS_AUTHOR
    region = NEWSPAPER_REGION_MAP[paper_name]
    docs = []
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, values_only=True):
        if has_author:
            date_val, title_val, content_val = row[4], row[1], row[2]
            url_val, polarity_val = row[5], row[8]
        else:
            date_val, title_val, content_val = row[3], row[1], row[2]
            url_val, polarity_val = row[4], row[7]
        if not date_val:
            continue
        docs.append({
            "create_date": convert_date(date_val),
            "site_type": "newspaper",
            "site_name": paper_name,
            "region": region,
            "title": safe_str(title_val) or None,
            "content": safe_str(content_val),
            "url": safe_str(url_val),
            "polarity": POLARITY_MAP.get(safe_str(polarity_val), "0"),
        })
    return docs


def build_trend_data(all_docs: list[dict]) -> list[dict]:
    counts = defaultdict(int)
    for doc in all_docs:
        day = doc["create_date"][:8]
        counts[(day, doc["site_type"])] += 1

    trend = []
    for (day, site_type), count in sorted(counts.items()):
        trend.append({
            "create_date": day,
            "site_type": site_type,
            "doc_count": count,
        })
    return trend


def main():
    os.makedirs(DATA_DIR, exist_ok=True)

    excel_dir = BASE_DIR
    all_docs = []

    all_docs.extend(load_twitter(str(excel_dir / "X(트위터).xlsx")))
    all_docs.extend(load_media(str(excel_dir / "매스미디어.xlsx")))
    all_docs.extend(load_youtube(str(excel_dir / "유튜브 스크립트.xlsx")))
    all_docs.extend(load_community(str(excel_dir / "커뮤니티.xlsx")))

    all_docs.sort(key=lambda x: x["create_date"], reverse=True)

    docs_path = DATA_DIR / "all_documents.json"
    with open(docs_path, "w", encoding="utf-8") as f:
        json.dump(all_docs, f, ensure_ascii=False, indent=2)

    trend = build_trend_data(all_docs)
    trend_path = DATA_DIR / "trend_data.json"
    with open(trend_path, "w", encoding="utf-8") as f:
        json.dump(trend, f, ensure_ascii=False, indent=2)

    print(f"Documents: {len(all_docs)} items -> {docs_path}")
    print(f"Trend: {len(trend)} items -> {trend_path}")

    by_type = defaultdict(int)
    for doc in all_docs:
        by_type[doc["site_type"]] += 1
    for st, cnt in sorted(by_type.items()):
        print(f"  {st}: {cnt}")

    # 신문 데이터 변환
    newspaper_docs = []
    for paper_name in NEWSPAPER_REGION_MAP:
        filepath = excel_dir / f"{paper_name}.xlsx"
        if filepath.exists():
            docs = load_newspaper(str(filepath), paper_name)
            newspaper_docs.extend(docs)
            print(f"  newspaper/{paper_name}: {len(docs)}")
        else:
            print(f"  newspaper/{paper_name}: SKIPPED (file not found)")

    newspaper_docs.sort(key=lambda x: x["create_date"], reverse=True)

    newspaper_path = DATA_DIR / "newspaper_articles.json"
    with open(newspaper_path, "w", encoding="utf-8") as f:
        json.dump(newspaper_docs, f, ensure_ascii=False, indent=2)

    print(f"Newspapers: {len(newspaper_docs)} items -> {newspaper_path}")


if __name__ == "__main__":
    main()
